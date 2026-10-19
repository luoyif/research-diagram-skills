#!/usr/bin/env python3
"""Research Diagram Skills, authored for 小罗. MIT; third-party code is separate.
No model/network calls. A vision-capable agent authors Scene v1 from a reference;
this tool validates and builds native Draw.io/PPTX and vector SVG from that scene.
"""
from __future__ import annotations
import argparse, copy, hashlib, json, math, os, platform, re, shutil, subprocess
import sys, tempfile, time, xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any

VERSION = '1.0.0'
SCENE_KINDS = {'group', 'rect', 'ellipse', 'text'}

def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def save_json(path: Path, value: Any) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')

def load(path: Path) -> dict:
    if path.stat().st_size > 10_000_000:
        raise ValueError('Scene exceeds 10 MB; use multiple pages.')
    return json.loads(path.read_text(encoding='utf-8'))

def finite(x: Any) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool) and math.isfinite(x)

def bounds(scene: dict) -> dict[str, tuple]:
    nodes = {n['id']: n for n in scene['nodes']}; cache = {}; visiting = set()
    def resolve(i: str) -> tuple:
        if i == 'root': return (0., 0., scene['width'], scene['height'])
        if i in cache: return cache[i]
        if i in visiting: raise ValueError('Cyclic containment: '+i)
        visiting.add(i); n=nodes[i]; p=resolve(n.get('parent','root'))
        cache[i]=(p[0]+n['x'],p[1]+n['y'],n['w'],n['h']); visiting.remove(i)
        return cache[i]
    for i in nodes: resolve(i)
    return cache

def edge_points(edge: dict, bb: dict) -> list[tuple[float,float]]:
    a,b=bb[edge['source']],bb[edge['target']]
    p=edge.get('source_port',[1,.5]); q=edge.get('target_port',[0,.5])
    return [(a[0]+a[2]*p[0],a[1]+a[3]*p[1])]+[tuple(x) for x in edge.get('waypoints',[])]+[(b[0]+b[2]*q[0],b[1]+b[3]*q[1])]

def ancestors(i: str, nodes: dict) -> set:
    out=set()
    while i in nodes:
        i=nodes[i].get('parent','root'); out.add(i)
    return out

def segment_inside(a, b, rect, eps=.5) -> bool:
    """Liang-Barsky clipping against a shrunken rectangle (boundary contact is OK)."""
    x,y,w,h=rect; xmin,ymin,xmax,ymax=x+eps,y+eps,x+w-eps,y+h-eps
    if xmin>=xmax or ymin>=ymax:return False
    dx,dy=b[0]-a[0],b[1]-a[1]; lo,hi=0.,1.
    for p,q in [(-dx,a[0]-xmin),(dx,xmax-a[0]),(-dy,a[1]-ymin),(dy,ymax-a[1])]:
        if abs(p)<1e-12:
            if q<0:return False
        else:
            t=q/p
            if p<0:lo=max(lo,t)
            else:hi=min(hi,t)
            if lo>hi:return False
    return hi>lo

def crossing(a,b,c,d) -> bool:
    def orient(p,q,r):return (q[0]-p[0])*(r[1]-p[1])-(q[1]-p[1])*(r[0]-p[0])
    return orient(a,b,c)*orient(a,b,d)<-1e-8 and orient(c,d,a)*orient(c,d,b)<-1e-8

def collinear_overlap(a,b,c,d) -> bool:
    dx,dy=b[0]-a[0],b[1]-a[1]
    if abs(dx)+abs(dy)<1e-9:return False
    if abs(dx*(c[1]-a[1])-dy*(c[0]-a[0]))>1e-7 or abs(dx*(d[1]-a[1])-dy*(d[0]-a[0]))>1e-7:return False
    axis=0 if abs(dx)>=abs(dy) else 1
    return min(max(a[axis],b[axis]),max(c[axis],d[axis]))-max(min(a[axis],b[axis]),min(c[axis],d[axis]))>1

def contract(scene: dict) -> dict:
    return {'schema_version':1,'text':{n['id']:n.get('text','') for n in scene['nodes'] if n.get('text')},
            'relations':{e['id']:{'source':e['source'],'target':e['target'],'arrow':e.get('arrow',True),'dashed':e.get('dashed',False)} for e in scene.get('edges',[])}}

def validate(scene: dict, frozen: dict|None=None) -> dict:
    errors=[]; warnings=[]
    if not isinstance(scene,dict):return {'status':'FAIL','errors':[{'code':'SCHEMA','detail':'Scene must be an object'}],'warnings':[]}
    def err(code,detail):errors.append({'code':code,'detail':detail})
    if scene.get('schema_version')!=1:err('SCHEMA','Expected schema_version=1')
    for k in ('width','height'):
        if not finite(scene.get(k)) or not 100<=scene[k]<=5000:err('CANVAS',k)
    sizes=scene.get('font_sizes',[26,21])
    if not isinstance(sizes,list) or len(sizes)!=2 or any(not finite(x) or not 8<=x<=100 for x in sizes) or sizes[0]==sizes[1]:err('FONTS','Exactly two finite font sizes required')
    font=scene.get('font_family','Songti SC')
    if not isinstance(font,str) or not font.strip() or any(ch in font for ch in ';<>\n\r'):err('FONT_NAME','Invalid font family')
    ns=scene.get('nodes',[]); es=scene.get('edges',[])
    if not isinstance(ns,list) or not ns:err('NODES','Non-empty node list required')
    if not isinstance(es,list):err('EDGES','Edge list required')
    if errors:return {'status':'FAIL','errors':errors,'warnings':warnings}
    if any(not isinstance(n,dict) for n in ns+es):return {'status':'FAIL','errors':[{'code':'SCHEMA','detail':'Nodes and edges must be objects'}],'warnings':[]}
    ids=[n.get('id') for n in ns]+[e.get('id') for e in es]
    if any(not isinstance(i,str) or not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_.-]*',i) or i in ['root','0','1'] for i in ids):err('IDS','Use unique semantic ASCII IDs; root,0,1 reserved')
    if errors:return {'status':'FAIL','errors':errors,'warnings':warnings}
    if len(set(ids))!=len(ids):err('DUPLICATE_ID','IDs must be unique')
    nodes={n.get('id'):n for n in ns}
    for n in ns:
        i=n['id']; parent=n.get('parent','root')
        if n.get('kind') not in SCENE_KINDS:err('UNSUPPORTED_NODE',i)
        if any(not finite(n.get(k)) for k in ('x','y','w','h')):err('GEOMETRY',i);continue
        if n['w']<=0 or n['h']<=0:err('GEOMETRY_SIZE',i)
        if parent!='root' and (parent not in nodes or nodes[parent].get('kind')!='group'):err('PARENT',i)
        if not isinstance(n.get('text',''),str):err('TEXT_TYPE',i)
        if n.get('text') and n.get('font_size',sizes[1]) not in sizes:err('FONT_COUNT',i)
        if n.get('kind')=='group' and n.get('text'):err('GROUP_TEXT','Put labels in native child text nodes: '+i)
        if n.get('role') in ['background','decoration'] and n.get('text'):err('ROLE_TEXT','Text-bearing nodes cannot bypass overlap checks: '+i)
        for color in ('fill','stroke','color'):
            v=n.get(color,'none')
            if v!='none' and not re.fullmatch(r'#[0-9A-Fa-f]{6}',v):err('COLOR',i+':'+color)
            elif scene.get('monochrome',True) and v!='none' and len({v[1:3].lower(),v[3:5].lower(),v[5:7].lower()})!=1:err('MONOCHROME',i+':'+color)
    for e in es:
        if e.get('source') not in nodes or e.get('target') not in nodes:err('MISSING_ENDPOINT',e.get('id'));continue
        if e['source']==e['target']:err('SELF_LOOP','Model a self-loop with distinct ports/nodes explicitly: '+e['id'])
        for k,default in [('source_port',[1,.5]),('target_port',[0,.5])]:
            p=e.get(k,default)
            if not isinstance(p,list) or len(p)!=2 or any(not finite(x) or not 0<=x<=1 for x in p) or not any(x in [0,1] for x in p):err('PORT',e['id']+':'+k)
        if not isinstance(e.get('waypoints',[]),list):err('WAYPOINT',e['id']);continue
        for p in e.get('waypoints',[]):
            if not isinstance(p,list) or len(p)!=2 or any(not finite(v) for v in p):err('WAYPOINT',e['id'])
    if errors:return {'status':'FAIL','errors':errors,'warnings':warnings}
    try:bb=bounds(scene)
    except (ValueError,KeyError) as ex:return {'status':'FAIL','errors':[{'code':'CONTAINMENT','detail':str(ex)}],'warnings':warnings}
    for n in ns:
        x,y,w,h=bb[n['id']]; parent=n.get('parent','root'); px,py,pw,ph=(0,0,scene['width'],scene['height']) if parent=='root' else bb[parent]
        if x<px-.01 or y<py-.01 or x+w>px+pw+.01 or y+h>py+ph+.01:err('OUT_OF_BOUNDS',n['id'])
        if n.get('text'):
            lines=n['text'].split('\n'); size=n.get('font_size',sizes[1])
            if len(lines)*size*1.22>h+.1:err('TEXT_HEIGHT',n['id'])
    # Intentional containment/backgrounds are excluded, not arbitrary overlapping siblings.
    solids=[n for n in ns if n['kind'] not in ['group'] and n.get('role') not in ['background','decoration']]
    for k,a in enumerate(solids):
        ax,ay,aw,ah=bb[a['id']]
        for b in solids[k+1:]:
            if a['id'] in ancestors(b['id'],nodes) or b['id'] in ancestors(a['id'],nodes):continue
            bx,by,bw,bh=bb[b['id']]
            if min(ax+aw,bx+bw)-max(ax,bx)>1 and min(ay+ah,by+bh)-max(ay,by)>1:err('SIBLING_OVERLAP',a['id']+' / '+b['id'])
    for e in es:
        pts=edge_points(e,bb); skip={e['source'],e['target']}|ancestors(e['source'],nodes)|ancestors(e['target'],nodes)
        for n in solids:
            if n['id'] in skip or e['source'] in ancestors(n['id'],nodes) or e['target'] in ancestors(n['id'],nodes):continue
            if any(segment_inside(a,b,bb[n['id']]) for a,b in zip(pts,pts[1:])):err('EDGE_THROUGH_NODE',e['id']+' / '+n['id'])
        if math.dist(pts[-2],pts[-1])<20:warnings.append({'code':'SHORT_ARROW_SEGMENT','detail':e['id']})
        for point in pts:
            if not (0<=point[0]<=scene['width'] and 0<=point[1]<=scene['height']):err('EDGE_OUT_OF_BOUNDS',e['id'])
    for k,a in enumerate(es):
        for b in es[k+1:]:
            pa,pb=edge_points(a,bb),edge_points(b,bb)
            if any(collinear_overlap(x,y,z,t) for x,y in zip(pa,pa[1:]) for z,t in zip(pb,pb[1:])):
                if a.get('shared_bus') and a.get('shared_bus')==b.get('shared_bus'):warnings.append({'code':'DECLARED_SHARED_BUS','detail':a['id']+' / '+b['id']})
                else:err('EDGE_STACKED',a['id']+' / '+b['id'])
            if any(crossing(x,y,z,t) for x,y in zip(pa,pa[1:]) for z,t in zip(pb,pb[1:])):
                if a.get('allow_crossing') and b.get('allow_crossing'):warnings.append({'code':'DECLARED_CROSSING','detail':a['id']+' / '+b['id']})
                else:err('EDGE_CROSSING',a['id']+' / '+b['id'])
    if frozen is not None and contract(scene)!=frozen:err('CONTRACT_DRIFT','Frozen text or relationship direction changed; review explicitly, do not refreeze silently')
    return {'status':'FAIL' if errors else 'PASS','errors':errors,'warnings':warnings,'node_count':len(ns),'edge_count':len(es),'visual_status':'NOT_REVIEWED','limitation':'Geometry checks do not prove reference fidelity, renderer text metrics, or scientific correctness.'}

def ordered(scene: dict) -> list[dict]:
    out=[]
    def walk(parent):
        children=[n for n in scene['nodes'] if n.get('parent','root')==parent]
        children.sort(key=lambda n:0 if n.get('role')=='background' else 1)
        for n in children:
            out.append(n)
            if n['kind']=='group':walk(n['id'])
    walk('root');return out

def drawio(scene: dict) -> bytes:
    mf=ET.Element('mxfile',host='research-diagram-skills',version=VERSION,author=scene.get('author','小罗'))
    page=ET.SubElement(mf,'diagram',id='research-page',name=scene.get('title','Diagram'))
    model=ET.SubElement(page,'mxGraphModel',grid='1',gridSize='10',page='1',pageScale='1',pageWidth=str(scene['width']),pageHeight=str(scene['height']),math='0',shadow='0')
    root=ET.SubElement(model,'root');ET.SubElement(root,'mxCell',id='0');ET.SubElement(root,'mxCell',id='1',parent='0')
    font=scene.get('font_family','Songti SC')
    for n in ordered(scene):
        k=n['kind'];styles={'html':'0','whiteSpace':'wrap','fontFamily':font,'fontSize':str(n.get('font_size',scene['font_sizes'][1])),'fontColor':n.get('color','#111111'),'fontStyle':'1' if n.get('bold') else '0','spacing':'0','align':n.get('align','center'),'verticalAlign':'middle','strokeWidth':str(n.get('stroke_width',1.5))}
        first=''
        if k=='group':first='group;';styles.update(connectable='1')
        elif k=='text':first='text;';styles.update(fillColor='none',strokeColor='none',connectable='0')
        else:
            if k=='ellipse':first='ellipse;'
            styles.update(fillColor=n.get('fill','#FFFFFF'),strokeColor=n.get('stroke','#333333'),rounded='1' if n.get('rounded',False) else '0',arcSize='8')
        cell=ET.SubElement(root,'mxCell',id=n['id'],value=n.get('text',''),style=first+';'.join(f'{a}={b}' for a,b in styles.items())+';',vertex='1',parent='1' if n.get('parent','root')=='root' else n['parent'])
        ET.SubElement(cell,'mxGeometry',x=str(n['x']),y=str(n['y']),width=str(n['w']),height=str(n['h']),attrib={'as':'geometry'})
    for e in scene.get('edges',[]):
        p,q=e.get('source_port',[1,.5]),e.get('target_port',[0,.5])
        st=f'edgeStyle=none;noEdgeStyle=1;html=0;rounded=0;strokeColor=#333333;strokeWidth=2;endArrow={"block" if e.get("arrow",True) else "none"};endFill=1;dashed={int(e.get("dashed",False))};exitX={p[0]};exitY={p[1]};entryX={q[0]};entryY={q[1]};exitPerimeter=0;entryPerimeter=0;'
        c=ET.SubElement(root,'mxCell',id=e['id'],value='',style=st,edge='1',parent='1',source=e['source'],target=e['target'])
        g=ET.SubElement(c,'mxGeometry',relative='1',attrib={'as':'geometry'})
        if e.get('waypoints'):
            ar=ET.SubElement(g,'Array',attrib={'as':'points'})
            for x,y in e['waypoints']:ET.SubElement(ar,'mxPoint',x=str(x),y=str(y))
    return ET.tostring(mf,encoding='utf-8',xml_declaration=True)

def vector_svg(scene: dict, model: bytes) -> bytes:
    ns='http://www.w3.org/2000/svg';ET.register_namespace('',ns);tag=lambda x:'{'+ns+'}'+x
    r=ET.Element(tag('svg'),width=str(scene['width']),height=str(scene['height']),viewBox=f"0 0 {scene['width']} {scene['height']}",content=model.decode())
    ET.SubElement(r,tag('rect'),x='0',y='0',width='100%',height='100%',fill='#FFFFFF')
    defs=ET.SubElement(r,tag('defs'));marker=ET.SubElement(defs,tag('marker'),id='arrow',viewBox='0 0 10 10',refX='9',refY='5',markerWidth='6',markerHeight='6',orient='auto');ET.SubElement(marker,tag('path'),d='M 0 0 L 10 5 L 0 10 z',fill='#333333')
    bb=bounds(scene);by={'root':r}
    for n in ordered(scene):
        parent=by[n.get('parent','root')];k=n['kind'];x,y,w,h=[n[v] for v in ['x','y','w','h']]
        if k=='group':by[n['id']]=ET.SubElement(parent,tag('g'),id=n['id'],transform=f'translate({x},{y})');continue
        if k in ['rect','ellipse']:
            attrs={'fill':n.get('fill','#FFFFFF'),'stroke':n.get('stroke','#333333'),'stroke-width':str(n.get('stroke_width',1.5))}
            if k=='rect':ET.SubElement(parent,tag('rect'),id=n['id'],x=str(x),y=str(y),width=str(w),height=str(h),rx='6' if n.get('rounded') else '0',attrib=attrs)
            else:ET.SubElement(parent,tag('ellipse'),id=n['id'],cx=str(x+w/2),cy=str(y+h/2),rx=str(w/2),ry=str(h/2),attrib=attrs)
        if n.get('text'):
            size=n.get('font_size',scene['font_sizes'][1]);lines=n['text'].split('\n');lh=size*1.22
            tx=x if n.get('align')=='left' else x+w/2
            t=ET.SubElement(parent,tag('text'),id=n['id']+'-label',x=str(tx),fill=n.get('color','#111111'),attrib={'font-family':scene.get('font_family','Songti SC'),'font-size':str(size),'font-weight':'bold' if n.get('bold') else 'normal','text-anchor':'start' if n.get('align')=='left' else 'middle'})
            for i,line in enumerate(lines):ET.SubElement(t,tag('tspan'),x=str(tx),y=str(y+(h-len(lines)*lh)/2+size*.95+i*lh)).text=line
    for e in scene.get('edges',[]):
        pts=edge_points(e,bb);attrs={'fill':'none','stroke':'#333333','stroke-width':'2'}
        if e.get('arrow',True):attrs['marker-end']='url(#arrow)'
        if e.get('dashed'):attrs['stroke-dasharray']='8 6'
        ET.SubElement(r,tag('path'),id=e['id'],d='M '+' L '.join(f'{x} {y}' for x,y in pts),attrib=attrs)
    return ET.tostring(r,encoding='utf-8',xml_declaration=True)

def native_ppt(scene: dict, destination: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR,MSO_AUTO_SIZE,PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    prs=Presentation();prs.slide_width=Inches(scene['width']/96);prs.slide_height=Inches(scene['height']/96)
    prs.core_properties.author=scene.get('author','小罗');prs.core_properties.last_modified_by=scene.get('author','小罗');prs.core_properties.title=scene.get('title','Research Diagram')
    slide=prs.slides.add_slide(prs.slide_layouts[6]);bb=bounds(scene);shapes={};em=lambda v:Inches(v/96)
    for n in ordered(scene):
        k=n['kind'];x,y,w,h=bb[n['id']]
        if k=='group':continue
        if k=='text':shape=slide.shapes.add_textbox(em(x),em(y),em(w),em(h))
        else:
            typ=MSO_SHAPE.OVAL if k=='ellipse' else (MSO_SHAPE.ROUNDED_RECTANGLE if n.get('rounded') else MSO_SHAPE.RECTANGLE)
            shape=slide.shapes.add_shape(typ,em(x),em(y),em(w),em(h))
            if n.get('rounded') and k=='rect':shape.adjustments[0]=.06
            for field,obj in [('fill',shape.fill),('stroke',shape.line.fill)]:
                col=n.get(field,'#FFFFFF' if field=='fill' else '#333333')
                if col=='none':obj.background()
                else:obj.solid();obj.fore_color.rgb=RGBColor.from_string(col[1:])
            shape.line.width=Pt(n.get('stroke_width',1.5)*.75)
        shape.name=n['id'];shapes[n['id']]=shape
        if n.get('text'):
            tf=shape.text_frame;tf.clear();tf.word_wrap=False;tf.auto_size=MSO_AUTO_SIZE.NONE;tf.vertical_anchor=MSO_ANCHOR.MIDDLE;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
            size=n.get('font_size',scene['font_sizes'][1])*.75
            for i,line in enumerate(n['text'].split('\n')):
                p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.alignment=PP_ALIGN.LEFT if n.get('align')=='left' else PP_ALIGN.CENTER;p.space_before=p.space_after=Pt(0);p.line_spacing=Pt(size*1.22)
                run=p.add_run();run.text=line;run.font.name=scene.get('font_family','Songti SC');run.font.size=Pt(size);run.font.bold=n.get('bold',False);run.font.color.rgb=RGBColor.from_string(n.get('color','#111111')[1:])
                rp=run._r.get_or_add_rPr()
                for field in ['a:ea','a:cs']:z=OxmlElement(field);z.set('typeface',scene.get('font_family','Songti SC'));rp.append(z)
    for n in reversed(ordered(scene)):
        if n['kind']=='group':
            children=[shapes[c['id']] for c in scene['nodes'] if c.get('parent','root')==n['id']]
            if children:g=slide.shapes.add_group_shape(children);g.name=n['id'];shapes[n['id']]=g
    for e in scene.get('edges',[]):
        pts=edge_points(e,bb)
        for i,(a,b) in enumerate(zip(pts,pts[1:])):
            c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,em(a[0]),em(a[1]),em(b[0]),em(b[1]));c.name=e['id']+'__segment_'+str(i);c.line.color.rgb=RGBColor(51,51,51);c.line.width=Pt(1.5);ln=c._element.spPr.get_or_add_ln()
            if e.get('dashed'):ETX=OxmlElement('a:prstDash');ETX.set('val','dash');ln.append(ETX)
            if i==len(pts)-2 and e.get('arrow',True):ETX=OxmlElement('a:tailEnd');ETX.set('type','triangle');ln.append(ETX)
    slide.notes_slide.notes_text_frame.text='Author: '+scene.get('author','小罗')+'\nNative editable text, shapes and connectors. Multi-segment PowerPoint connectors require rechecking after moving nodes. Draw.io remains the preferred topology-editing master.'
    for n in slide._element.xpath('.//a:effectRef'):n.set('idx','0')
    prs.save(destination)

def discover_drawio(explicit: str|None=None) -> str|None:
    candidates=[explicit,os.getenv('DRAWIO_PATH'),shutil.which('drawio'),shutil.which('draw.io'),'/Applications/draw.io.app/Contents/MacOS/draw.io']
    if os.name=='nt':candidates += [str(Path(os.getenv('LOCALAPPDATA',''))/'Programs'/'draw.io'/'draw.io.exe'),str(Path(os.getenv('ProgramFiles','C:/Program Files'))/'draw.io'/'draw.io.exe')]
    return next((str(Path(p).resolve()) for p in candidates if p and Path(p).is_file()),None)

def build(scene_path: Path, out: Path, mode: str, frozen_path: Path|None=None) -> dict:
    scene=load(scene_path);report=validate(scene,load(frozen_path) if frozen_path else None)
    if report['status']!='PASS':raise ValueError(json.dumps(report,ensure_ascii=False))
    # Never overwrite a previous result. The input scene remains unchanged.
    out.mkdir(parents=True,exist_ok=False)
    shutil.copy2(scene_path,out/'scene.json');save_json(out/'geometry-check.json',report)
    model=drawio(scene)
    if mode in ['drawio','both']:(out/'diagram.drawio').write_bytes(model)
    if mode in ['pptx','both']:native_ppt(scene,out/'diagram.pptx')
    (out/'diagram.svg').write_bytes(vector_svg(scene,model))
    state={'version':VERSION,'mode':mode,'geometry':'PASS','visual_review':'PENDING','render':{},'build_hashes':{p.name:sha(p) for p in out.iterdir() if p.is_file()},'font_requested':scene.get('font_family','Songti SC'),'created_at_utc':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())}
    save_json(out/'build-state.json',state);return state

def render(out: Path, mode: str, drawio_bin: str|None, width: int) -> dict:
    from PIL import Image
    if not 500<=width<=20000:raise ValueError('Width must be 500..20000 pixels')
    state=load(out/'build-state.json')
    # Only direct CLI invocation, no shell string assembled from diagram content.
    if mode=='drawio':
        binary=discover_drawio(drawio_bin)
        if not binary:raise ValueError('draw.io Desktop CLI not found; native source is available, render remains PENDING')
        src=out/'diagram.drawio';dest=out/'preview-drawio.png'
        if not src.exists():raise ValueError('Build drawio or both mode first')
        with tempfile.TemporaryDirectory(prefix='rds-render-') as tmp:
            target=Path(tmp)/'render.png';cmd=[binary,'-x','-f','png','--width',str(width),'-b','10','-o',str(target),str(src.resolve())]
            p=subprocess.run(cmd,capture_output=True,text=True,timeout=90)
            if p.returncode or not target.exists() or not target.stat().st_size:raise RuntimeError('Renderer failed: '+(p.stderr+p.stdout)[-1500:])
            with Image.open(target) as im:im.load();im.save(dest,dpi=(600,600));pixels=list(im.size)
        record={'engine':'draw.io Desktop','source_sha256':sha(src),'preview_sha256':sha(dest),'pixels':pixels,'dpi':600}
    else:
        soffice=shutil.which('soffice') or ('/Applications/LibreOffice.app/Contents/MacOS/soffice' if Path('/Applications/LibreOffice.app/Contents/MacOS/soffice').exists() else None)
        ppm=shutil.which('pdftoppm')
        if not soffice or not ppm:raise ValueError('PPT rendering needs LibreOffice + pdftoppm; do not substitute SVG PNG for PPT QA')
        src=out/'diagram.pptx';dest=out/'preview-pptx.png'
        with tempfile.TemporaryDirectory(prefix='rds-ppt-') as tmp:
            t=Path(tmp);cmd=[soffice,'-env:UserInstallation='+str((t/'profile').as_uri()),'--headless','--convert-to','pdf','--outdir',str(t),str(src.resolve())]
            p=subprocess.run(cmd,capture_output=True,text=True,timeout=90);pdf=t/'diagram.pdf'
            if p.returncode or not pdf.exists():raise RuntimeError('PPT renderer failed: '+(p.stdout+p.stderr)[-1500:])
            p=subprocess.run([ppm,'-png','-scale-to',str(width),'-singlefile',str(pdf),str(t/'page')],capture_output=True,text=True,timeout=90)
            if p.returncode or not (t/'page.png').exists():raise RuntimeError('PPT raster preview failed')
            with Image.open(t/'page.png') as im:im.load();im.save(dest,dpi=(600,600));pixels=list(im.size)
        record={'engine':'LibreOffice + pdftoppm','source_sha256':sha(src),'preview_sha256':sha(dest),'pixels':pixels,'dpi':600}
    state['render'][mode]=record;state['visual_review']='PENDING';state.pop('review',None);save_json(out/'build-state.json',state);return record

def review(out: Path, reviewer: str, notes: str, verdict: str) -> dict:
    state=load(out/'build-state.json')
    for name,h in state['build_hashes'].items():
        if not (out/name).exists() or sha(out/name)!=h:raise ValueError('Build artifact changed: '+name+'; rebuild into a fresh pass and rerender')
    expected={'drawio','pptx'} if state['mode']=='both' else {state['mode']}
    for mode in expected:
        rec=state['render'].get(mode)
        if not rec:raise ValueError('Missing actual '+mode+' preview; review cannot pass')
        if rec['source_sha256']!=sha(out/('diagram.drawio' if mode=='drawio' else 'diagram.pptx')) or rec['preview_sha256']!=sha(out/('preview-'+mode+'.png')):raise ValueError('Source or render changed; rerender before review')
    if not reviewer.strip() or len(notes.strip())<12:raise ValueError('Named reviewer and specific review notes required')
    state['visual_review']=verdict
    state['review']={'reviewer':reviewer,'notes':notes,'verdict':verdict,'scope':'Human/agent-declared visual inspection; not an automatic proof','artifacts':{p.name:sha(p) for p in out.iterdir() if p.is_file() and p.name!='build-state.json'}}
    save_json(out/'build-state.json',state);return state['review']

def main() -> int:
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='command',required=True)
    p=sub.add_parser('doctor');p.add_argument('--drawio')
    p=sub.add_parser('validate');p.add_argument('scene',type=Path);p.add_argument('--contract',type=Path)
    p=sub.add_parser('freeze');p.add_argument('scene',type=Path);p.add_argument('destination',type=Path)
    p=sub.add_parser('build');p.add_argument('scene',type=Path);p.add_argument('--out',required=True,type=Path);p.add_argument('--mode',choices=['drawio','pptx','both'],default='both');p.add_argument('--contract',type=Path)
    p=sub.add_parser('render');p.add_argument('out',type=Path);p.add_argument('--mode',choices=['drawio','pptx'],required=True);p.add_argument('--drawio');p.add_argument('--width',type=int,default=6400)
    p=sub.add_parser('review');p.add_argument('out',type=Path);p.add_argument('--reviewer',required=True);p.add_argument('--notes',required=True);p.add_argument('--verdict',choices=['PASS','FAIL'],required=True)
    args=parser.parse_args()
    try:
        if args.command=='doctor':
            import importlib.util
            data={'version':VERSION,'platform':platform.system(),'drawio':discover_drawio(args.drawio),'python_pptx':bool(importlib.util.find_spec('pptx')),'Pillow':bool(importlib.util.find_spec('PIL')),'soffice':shutil.which('soffice'),'pdftoppm':shutil.which('pdftoppm'),'font_note':'Fonts are not bundled. Select an installed family; substitute only with disclosure.'}
        elif args.command=='validate':
            data=validate(load(args.scene),load(args.contract) if args.contract else None);print(json.dumps(data,ensure_ascii=False,indent=2));return 0 if data['status']=='PASS' else 2
        elif args.command=='freeze':
            if args.destination.exists():raise ValueError('Contract exists: explicit human review needed before changing it')
            scene=load(args.scene);assert validate(scene)['status']=='PASS';data=contract(scene);save_json(args.destination,data)
        elif args.command=='build':data=build(args.scene,args.out,args.mode,args.contract)
        elif args.command=='render':data=render(args.out,args.mode,args.drawio,args.width)
        else:data=review(args.out,args.reviewer,args.notes,args.verdict)
        print(json.dumps(data,ensure_ascii=False,indent=2));return 0
    except (ValueError,KeyError,AssertionError,RuntimeError,ImportError,FileExistsError,subprocess.TimeoutExpired) as ex:
        print('ERROR: '+str(ex),file=sys.stderr);return 2

if __name__=='__main__':raise SystemExit(main())
